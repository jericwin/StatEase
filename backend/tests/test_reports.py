"""Compare planning, objectives → report, editable blocks with recomputation, and report exports."""

from __future__ import annotations

import io

import numpy as np
import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook

from app.reports.blocks import MANUAL_NOTE, resolve_chart, resolve_table

API = "/api/v1"


def formulation_csv() -> bytes:
    rng = np.random.default_rng(3)
    rows = []
    for f, shift in (("Formulation 1", 0), ("Formulation 2", 1.2)):
        for _ in range(25):
            rows.append({"Formulation": f, "Taste": int(np.clip(round(6 + shift + rng.normal(0, 1)), 1, 9)),
                         "Texture": int(np.clip(round(6.5 + rng.normal(0, 1)), 1, 9)),
                         "Overall Acceptability": round(float(6.2 + shift * 0.8 + rng.normal(0, 0.8)), 2)})
    return pd.DataFrame(rows).to_csv(index=False).encode()


@pytest.fixture(scope="module")
def setup(client):
    p = client.post(f"{API}/projects", json={"name": "Formulation study"}).json()
    up = client.post(f"{API}/datasets/uploads", files={"file": ("f.csv", formulation_csv(), "text/csv")}).json()
    levels = {v["column_name"]: v["measurement_level"] for v in up["variables"]}
    assert levels["Formulation"] in ("binary", "nominal")
    ds = client.post(f"{API}/datasets", json={"project_id": p["id"], "upload_id": up["upload"]["upload_id"],
                                             "variables": {"Taste": {"measurement_level": "numeric"},
                                                           "Texture": {"measurement_level": "numeric"}}}).json()
    return p, ds


def test_recommends_bundled_comparison(client, setup):
    _, ds = setup
    recs = client.get(f"{API}/analyses/recommendations", params={"dataset_id": ds["id"]}).json()["recommendations"]
    top = [r for r in recs if r["analysis_type"] == "compare_two_groups"]
    assert top and top[0]["title"] == "Compare Formulation 1 and Formulation 2"
    assert recs.index(top[0]) <= 1


def test_compare_plan(client, setup):
    _, ds = setup
    plan = client.post(f"{API}/analyses/compare/plan", json={"dataset_id": ds["id"], "mode": "groups",
                                                            "group": "Formulation", "measures": ["Taste", "Overall Acceptability"]}).json()
    assert plan["analysis_type"] == "compare_two_groups"
    assert plan["options"]["group_a"] == "Formulation 1"
    q = client.post(f"{API}/analyses/compare/plan", json={"dataset_id": ds["id"], "mode": "variables",
                                                         "var1": "Taste", "var2": "Texture"}).json()
    assert "question" in q  # paired vs relationship must be asked, not guessed
    paired = client.post(f"{API}/analyses/compare/plan", json={"dataset_id": ds["id"], "mode": "variables", "var1": "Taste",
                                                              "var2": "Texture", "relation": "paired"}).json()
    assert paired["analysis_type"] in ("paired_t", "wilcoxon")
    bad = client.post(f"{API}/analyses/compare/plan", json={"dataset_id": ds["id"], "mode": "groups", "group": "Taste",
                                                           "measures": ["Texture"]})
    assert bad.status_code == 422


def test_objectives_report_edit_recompute_export(client, setup):
    p, ds = setup
    text = ("1. Determine the acceptability of Formulation 1.\n2. Determine the acceptability of Formulation 2.\n"
            "3. Determine whether a significant difference exists between the two formulations.\n"
            "4. Determine the purple elephants.")
    objs = client.post(f"{API}/reports/objectives/parse", json={"dataset_id": ds["id"], "text": text}).json()["objectives"]
    assert [o["analysis_type"] for o in objs[:3]] == ["descriptives", "descriptives", "compare_two_groups"]
    assert objs[0]["filter"] == {"column": "Formulation", "values": ["Formulation 1"]}
    assert objs[2]["options"] == {"group_a": "Formulation 1", "group_b": "Formulation 2"}

    rep = client.post(f"{API}/reports", json={"project_id": p["id"], "dataset_id": ds["id"], "title": "Sensory Study",
                                             "objectives": objs[:3]}).json()
    tables = [b for b in rep["blocks"] if b["type"] == "table"]
    f1 = tables[0]
    assert "Formulation 1" in " ".join(f1["analysis"]["filter"]["values"])
    mean_f1 = f1["resolved"]["rows"][0]["mean"]

    # Replace Formulation 1 with Formulation 2: values must be recomputed, not relabelled.
    rep2 = client.post(f"{API}/reports/{rep['id']}/blocks/{f1['id']}/rebind",
                       json={"filter": {"column": "Formulation", "values": ["Formulation 2"]}}).json()
    f2 = next(b for b in rep2["blocks"] if b["id"] == f1["id"])
    assert f2["analysis"]["filter"]["values"] == ["Formulation 2"]
    assert f2["resolved"]["rows"][0]["mean"] != pytest.approx(mean_f1)
    src = pd.read_csv(io.BytesIO(formulation_csv()))
    var = f2["resolved"]["rows"][0]["variable"]
    assert f2["resolved"]["rows"][0]["mean"] == pytest.approx(src.loc[src.Formulation == "Formulation 2", var].mean())

    # Both formulations side by side.
    rep3 = client.post(f"{API}/reports/{rep['id']}/blocks/{f1['id']}/rebind",
                       json={"filter": {"column": "Formulation", "values": ["Formulation 1", "Formulation 2"]},
                             "variables": {**f2["analysis"]["variables"], "split_by": "Formulation"}}).json()
    both = next(b for b in rep3["blocks"] if b["id"] == f1["id"])
    assert {r["group"] for r in both["resolved"]["rows"]} == {"Formulation 1", "Formulation 2"}

    # Edits: rename, reorder, manual cell, hide a block, reorder blocks, then export.
    blocks = rep3["blocks"]
    tb = next(b for b in blocks if b["id"] == f1["id"])
    cols = [c["key"] for c in tb["source_table"]["columns"]]
    tb["title"], tb["title_edited"] = "Table A. Acceptability of both formulations", True
    tb["edits"] = {"columns": [{"key": cols[1]}, {"key": cols[0], "label": "Attribute"}],
                   "cells": {"0:n": 999}}
    chart = next(b for b in blocks if b["type"] == "chart" and b["include"])
    chart["include"] = False
    blocks.remove(tb)
    blocks.insert(1, tb)
    saved = client.put(f"{API}/reports/{rep['id']}", json={"blocks": blocks}).json()
    st = next(b for b in saved["blocks"] if b["id"] == f1["id"])
    assert [c["key"] for c in st["resolved"]["columns"]][:2] == [cols[1], cols[0]]
    assert st["resolved"]["columns"][1]["label"] == "Attribute"
    assert MANUAL_NOTE in st["resolved"]["notes"]

    r = client.post(f"{API}/exports", json={"project_id": p["id"], "format": "docx", "report_id": rep["id"]})
    assert r.status_code == 201, r.text
    doc = Document(io.BytesIO(client.get(r.json()["download_url"]).content))
    text_all = "\n".join(par.text for par in doc.paragraphs)
    assert "Table A. Acceptability of both formulations" in text_all
    assert "Objectives of the Study" in text_all
    assert any("999†" == c.text for t in doc.tables for row in t.rows for c in row.cells)
    order = [par.text for par in doc.paragraphs if par.text.startswith(("Objective 1", "Table A"))]
    assert order[0].startswith("Table A") or text_all.index("Table A") < text_all.index("Objective 2")

    r = client.post(f"{API}/exports", json={"project_id": p["id"], "format": "xlsx", "report_id": rep["id"],
                                            "include_data": True})
    wb = load_workbook(io.BytesIO(client.get(r.json()["download_url"]).content))
    assert {"Summary", "Descriptive Statistics", "Comparisons", "Data"} <= set(wb.sheetnames)
    if "Charts" in wb.sheetnames:
        assert wb["Charts"]._charts  # native, editable Excel charts
    comp = wb["Comparisons"]
    assert any(isinstance(c.value, float) for row in comp.iter_rows() for c in row)

    # Details off removes statistic/df/CI/effect columns from exports.
    client.put(f"{API}/reports/{rep['id']}", json={"options": {"include_details": False}})
    r = client.post(f"{API}/exports", json={"project_id": p["id"], "format": "docx", "report_id": rep["id"]})
    doc = Document(io.BytesIO(client.get(r.json()["download_url"]).content))
    headers = {c.text for t in doc.tables for c in t.rows[0].cells}
    assert "Statistic" not in headers and "p" in headers


def test_report_from_existing_results_and_library(client, setup):
    p, _ = setup
    rep = client.post(f"{API}/reports", json={"project_id": p["id"], "title": "All results"}).json()
    assert any(b["type"] == "table" for b in rep["blocks"])
    lib = client.get(f"{API}/reports/{rep['id']}/library").json()["items"]
    assert lib and lib[0]["tables"]


def test_public_demos_need_no_login(client, monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "auth_mode", "password")
    items = client.get(f"{API}/public/demos").json()["items"]
    assert [d["key"] for d in items] == ["fertilizer", "experiment", "survey"]
    pv = client.get(f"{API}/public/demos/fertilizer").json()
    main = next(t for t in pv["results"][0]["tables"] if t["id"] == "main")
    assert len(main["rows"]) == 4 and all(r["p"] < 0.05 for r in main["rows"])
    assert client.get(f"{API}/public/demos/fertilizer/sample.csv").status_code == 200


def test_block_edit_rules():
    src = {"id": "main", "kind": "main", "title": "T",
           "columns": [{"key": "a", "label": "A", "fmt": "text"}, {"key": "x", "label": "X", "fmt": "num", "group": "statistic"},
                       {"key": "p", "label": "p", "fmt": "p", "group": "p"}],
           "rows": [{"a": "r1", "x": 1.0, "p": 0.01}, {"a": "r2", "x": 2.0, "p": 0.2}], "notes": []}
    out = resolve_table(src, {"columns": [{"key": "p"}, {"key": "a", "label": "Row"}, {"key": "m_1", "label": "Note",
                                                                                      "manual": True}],
                              "hidden_rows": [1], "extra_rows": [{"a": "added"}]}, hide_details=True)
    assert [c["key"] for c in out["columns"]] == ["p", "a", "m_1"]  # x hidden as a detail column
    assert [r["a"] for r in out["rows"]] == ["r1", "added"]
    assert out["rows"][1]["_manual"] and MANUAL_NOTE in out["notes"]

    chart = {"id": "means", "type": "meanci", "title": "M", "x_label": "", "y_label": "Score", "caption": "",
             "data": {"groups": [{"name": "F1", "mean": 5, "lower": 4, "upper": 6},
                                 {"name": "F2", "mean": 7, "lower": 6, "upper": 8}]}}
    c = resolve_chart(chart, {"order": ["F2", "F1"], "labels": {"F1": "Formulation 1"}})
    assert [g["name"] for g in c["data"]["groups"]] == ["F2", "Formulation 1"]
    c = resolve_chart(chart, {"type": "means", "hidden": ["F1"]})
    assert c["type"] == "means" and c["data"]["categories"] == ["F2"] and c["data"]["series"][0]["values"] == [7]


def test_thesis_objectives_with_subitems(client):
    """The statement-of-the-problem format: numbered questions with 1.1/2.1 sub-items."""
    import numpy as np
    import pandas as pd

    rng = np.random.default_rng(9)
    rows = []
    for f, shift in (("Formulation A", 0.0), ("Formulation B", 0.8)):
        for _ in range(25):
            rows.append({"Respondent": f"R{len(rows) + 1}", "Age": int(rng.integers(18, 40)),
                         "Gender": rng.choice(["Male", "Female"]), "Formulation": f,
                         **{k: int(np.clip(round(6.5 + shift + rng.normal(0, 1)), 1, 9))
                            for k in ("Appearance", "Aroma", "Taste", "Texture", "Overall acceptability")}})
    csv = pd.DataFrame(rows).to_csv(index=False).encode()
    p = client.post(f"{API}/projects", json={"name": "Mochi"}).json()
    up = client.post(f"{API}/datasets/uploads", files={"file": ("mochi.csv", csv, "text/csv")}).json()
    ds = client.post(f"{API}/datasets", json={"project_id": p["id"], "upload_id": up["upload"]["upload_id"], "variables": {
        k: {"measurement_level": "numeric"} for k in ("Appearance", "Aroma", "Taste", "Texture", "Overall acceptability")}}).json()
    text = """1.    What is the demographic profile of the respondents in terms of:
1.1.    Age
1.2.    Gender
2.    What is the level of acceptability of the developed mochi using unripe banana flour in terms of:
2.1 Appearance
2.2 Aroma
2.3 Taste
2.4 Texture
2.5 Overall acceptability
3.    Which formulation of mochi using unripe banana flour is the most acceptable?
4.    Is there a significant difference in the level of acceptability between two formulations unripe banana as flour in producing mochi?"""
    objs = client.post(f"{API}/reports/objectives/parse", json={"dataset_id": ds["id"], "text": text}).json()["objectives"]
    assert len(objs) == 4
    assert objs[0]["subitems"] == ["Age", "Gender"]
    assert {a["analysis_type"] for a in objs[0]["analyses"]} == {"frequencies", "descriptives"}
    assert objs[1]["analyses"][0]["variables"]["split_by"] == "Formulation"
    assert len(objs[1]["analyses"][0]["variables"]["variables"]) == 5
    assert objs[2]["analysis_type"] == "rank_groups"
    assert objs[2]["options"]["decide_by"] == "Overall acceptability"
    assert objs[3]["analysis_type"] == "compare_two_groups"
    rep = client.post(f"{API}/reports", json={"project_id": p["id"], "dataset_id": ds["id"], "objectives": objs}).json()
    assert all(o["error"] is None for o in rep["objectives"]), rep["objectives"]
    rank = next(b for b in rep["blocks"] if b["type"] == "table" and b["source_table"]["title"].startswith("Ranking"))
    src = pd.read_csv(__import__("io").BytesIO(csv))
    best = src.groupby("Formulation")["Overall acceptability"].mean().idxmax()
    assert rank["resolved"]["rows"][0]["group"] == best
    # Only objective-related results: every analysis block belongs to an objective.
    assert all(b.get("objective_id") for b in rep["blocks"])


def test_chart_styles_render_for_word_and_excel(tmp_path):
    from openpyxl import Workbook, load_workbook

    from app.reports.blocks import chart_styles, resolve_chart
    from app.reports.charts import render
    from app.reports.xlsx_export import _native_chart

    means = {"id": "m", "type": "means", "title": "Mean growth", "x_label": "Measure", "y_label": "Mean", "caption": "",
             "note": None, "data": {"categories": ["Height", "Leaves"], "series": [
                 {"name": "Control", "values": [16.6, 6.8], "lower": [16.0, 6.3], "upper": [17.2, 7.3]},
                 {"name": "Vermicompost", "values": [21.6, 9.4], "lower": [20.8, 8.9], "upper": [22.4, 9.9]}]}}
    counts = {"id": "b", "type": "bar", "title": "Gender", "x_label": "Gender", "y_label": "Count", "caption": "",
              "note": None, "data": {"categories": ["Female", "Male"], "series": [{"name": "Count", "values": [33, 27]}]}}
    meanci = {"id": "c", "type": "meanci", "title": "Means", "x_label": "Group", "y_label": "Score", "caption": "",
              "note": None, "data": {"groups": [{"name": "A", "mean": 5.0, "lower": 4.5, "upper": 5.5, "n": 10},
                                                {"name": "B", "mean": 6.0, "lower": 5.4, "upper": 6.6, "n": 10}]}}
    # Means never offer stacking or pie; a single set of counts offers pie.
    assert chart_styles(means) == ["bar", "hbar", "line", "dot"]
    assert chart_styles(counts) == ["bar", "hbar", "line", "dot", "pie"]
    assert chart_styles({**means, "type": "histogram"}) == []
    assert resolve_chart(means, {"style": "pie"}).get("style") is None  # not allowed -> ignored
    assert resolve_chart(meanci, {"style": "bar"})["type"] == "means"  # point chart becomes bars with CIs
    assert resolve_chart(meanci, {"style": "dot"})["type"] == "meanci"

    wb = Workbook()
    ws = wb.active
    row = 1
    for spec in (means, counts, meanci):
        for style in chart_styles(spec):
            styled = resolve_chart(spec, {"style": style})
            assert render(styled)[:4] == b"\x89PNG"
            row = _native_chart(ws, styled, row) or row
    assert {type(c).__name__ for c in ws._charts} == {"BarChart", "LineChart", "PieChart"}
    wb.save(tmp_path / "c.xlsx")
    load_workbook(tmp_path / "c.xlsx")  # the workbook opens
