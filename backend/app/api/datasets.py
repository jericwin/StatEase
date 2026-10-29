from __future__ import annotations

import json
from typing import Any

import polars as pl
from fastapi import APIRouter, File, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select

from ..data import cleaning, importer, store
from ..data.inference import LEVELS
from ..deps import DB, CurrentUser, audit, get_dataset, get_project, get_version
from ..errors import AppError
from ..models import Dataset, Variable
from ..services import VAR_FIELDS, add_version, create_dataset, latest_version, refresh_profile
from ..storage import new_relative_path, remove, resolve, safe_filename

router = APIRouter(tags=["datasets"])


def variable_out(v: Variable, profile: dict[str, Any] | None = None) -> dict[str, Any]:
    p = (profile or {}).get(v.column_name, {})
    return {"id": v.id, "name": v.column_name, "label": v.variable_label, "level": v.measurement_level,
            "data_type": v.data_type, "inferred_level": v.inferred_level, "inference_note": v.inference_note,
            "category_labels": v.category_labels_json or {}, "category_order": v.category_order_json or [],
            "missing_codes": v.missing_value_codes_json or [], "unit": v.unit, "scale_group": v.scale_group,
            "profile": p}


def dataset_out(ds: Dataset, full: bool = False) -> dict[str, Any]:
    v = latest_version(ds)
    out = {"id": ds.id, "project_id": ds.project_id, "name": ds.name, "source_filename": ds.source_filename,
           "format": ds.format, "row_count": v.row_count, "column_count": ds.column_count,
           "created_at": ds.created_at.isoformat(), "lineage": ds.lineage_json or {},
           "parent_version_id": ds.parent_version_id,
           "version": {"id": v.id, "number": v.version_number, "note": v.note, "sha256": v.content_hash,
                       "created_at": v.created_at.isoformat()},
           "versions": [{"id": x.id, "number": x.version_number, "note": x.note, "rows": x.row_count,
                         "created_at": x.created_at.isoformat()} for x in ds.versions]}
    if full:
        prof = {c["name"]: c for c in (v.profile_json or {}).get("variables", [])}
        out["variables"] = [variable_out(x, prof) for x in v.variables]
        out["profile"] = {k: val for k, val in (v.profile_json or {}).items() if k != "variables"}
    return out


# --------------------------------------------------------------------------- import wizard


@router.post("/datasets/uploads", status_code=201)
def upload(user: CurrentUser, file: UploadFile = File(...)):
    meta = importer.stage_upload(file.filename or "upload", file.file)
    result = importer.preview(meta["upload_id"], {})
    return result


class PreviewIn(BaseModel):
    options: dict[str, Any] = Field(default_factory=dict)


@router.post("/datasets/uploads/{upload_id}/preview")
def preview(upload_id: str, body: PreviewIn, user: CurrentUser):
    return importer.preview(upload_id, body.options)


@router.delete("/datasets/uploads/{upload_id}", status_code=204)
def discard(upload_id: str, user: CurrentUser):
    importer.discard_upload(upload_id)


class CommitIn(BaseModel):
    project_id: int
    upload_id: str
    name: str | None = Field(None, max_length=200)
    options: dict[str, Any] = Field(default_factory=dict)
    # column -> {measurement_level, variable_label, category_order, ...}
    variables: dict[str, dict[str, Any]] = Field(default_factory=dict)


@router.post("/datasets", status_code=201)
def commit_import(body: CommitIn, user: CurrentUser, db: DB):
    project = get_project(db, user, body.project_id)
    for col, o in body.variables.items():
        if o.get("measurement_level") and o["measurement_level"] not in LEVELS:
            raise AppError("invalid_level", f"Invalid measurement level for '{col}'.", 422)
    out = importer.materialize(body.upload_id, body.options)
    meta = out["meta"]
    name = body.name or meta["filename"].rsplit(".", 1)[0]
    if out["detected"].get("sheet") and len(out["detected"].get("sheets", [])) > 1:
        name = body.name or f"{name} — {out['detected']['sheet']}"
    ds = create_dataset(db, project, name=name, data_path=out["data_path"], fmt=meta["format"],
                        source_filename=meta["filename"], original_path=out["original_path"],
                        overrides=body.variables,
                        note="Imported" + (f" (sheet '{out['detected']['sheet']}')" if out["detected"].get("sheet") else ""))
    ds.lineage_json = {"import": {"options": body.options, "issues": out["issues"], "source_sha256": meta["sha256"]}}
    audit(db, user, "dataset.import", project.id, dataset=ds.id, rows=out["row_count"], file=meta["filename"])
    db.commit()
    return dataset_out(ds, full=True)


# --------------------------------------------------------------------------- browsing


@router.get("/projects/{project_id}/datasets")
def list_datasets(project_id: int, user: CurrentUser, db: DB):
    get_project(db, user, project_id)
    rows = db.scalars(select(Dataset).where(Dataset.project_id == project_id).order_by(Dataset.created_at)).all()
    return {"items": [dataset_out(d) for d in rows], "total": len(rows)}


@router.get("/datasets/{dataset_id}")
def read_dataset(dataset_id: int, user: CurrentUser, db: DB):
    return dataset_out(get_dataset(db, user, dataset_id), full=True)


class DatasetPatch(BaseModel):
    name: str = Field(min_length=1, max_length=200)


@router.patch("/datasets/{dataset_id}")
def rename_dataset(dataset_id: int, body: DatasetPatch, user: CurrentUser, db: DB):
    ds = get_dataset(db, user, dataset_id)
    ds.name = body.name.strip()
    db.commit()
    return dataset_out(ds)


@router.delete("/datasets/{dataset_id}", status_code=204)
def delete_dataset(dataset_id: int, user: CurrentUser, db: DB):
    ds = get_dataset(db, user, dataset_id)
    paths = [v.data_path for v in ds.versions] + ([ds.storage_path] if ds.storage_path else [])
    audit(db, user, "dataset.delete", ds.project_id, dataset=ds.id, name=ds.name)
    db.delete(ds)
    db.commit()
    for rel in paths:
        remove(rel)


def _version(db, user, dataset_id: int, version_id: int | None):
    ds = get_dataset(db, user, dataset_id)
    if version_id is None:
        return latest_version(ds)
    v = get_version(db, user, version_id)
    if v.dataset_id != ds.id:
        raise AppError("invalid_version", "Version does not belong to this dataset.", 422)
    return v


@router.get("/datasets/{dataset_id}/preview")
def page(dataset_id: int, user: CurrentUser, db: DB, offset: int = Query(0, ge=0), limit: int = Query(100, ge=1, le=500),
         sort: str | None = None, desc: bool = False, search: str | None = None, filters: str | None = None,
         columns: str | None = None, version_id: int | None = None):
    v = _version(db, user, dataset_id, version_id)
    try:
        flt = json.loads(filters) if filters else []
    except json.JSONDecodeError as exc:
        raise AppError("invalid_filter", "Filters must be JSON.", 422) from exc
    return store.page(v, offset=offset, limit=limit, sort=sort, descending=desc, search=search or None, filters=flt,
                      columns=columns.split(",") if columns else None)


@router.get("/datasets/{dataset_id}/profile")
def profile(dataset_id: int, user: CurrentUser, db: DB, version_id: int | None = None):
    v = _version(db, user, dataset_id, version_id)
    return v.profile_json


@router.get("/datasets/{dataset_id}/variables")
def variables(dataset_id: int, user: CurrentUser, db: DB, version_id: int | None = None):
    v = _version(db, user, dataset_id, version_id)
    prof = {c["name"]: c for c in (v.profile_json or {}).get("variables", [])}
    return {"items": [variable_out(x, prof) for x in v.variables], "version_id": v.id}


@router.get("/datasets/{dataset_id}/variables/{name}/values")
def values(dataset_id: int, name: str, user: CurrentUser, db: DB, version_id: int | None = None):
    v = _version(db, user, dataset_id, version_id)
    if name not in {x.column_name for x in v.variables}:
        raise AppError("unknown_variable", f"Variable '{name}' not found.", 404)
    return {"items": store.column_values(v, name)}


class VariablePatch(BaseModel):
    label: str | None = Field(None, max_length=255)
    level: str | None = None
    category_labels: dict[str, str] | None = None
    category_order: list[Any] | None = None
    missing_codes: list[Any] | None = None
    unit: str | None = Field(None, max_length=50)
    scale_group: str | None = Field(None, max_length=120)


class VariablesPatch(BaseModel):
    variables: dict[str, VariablePatch]


@router.patch("/datasets/{dataset_id}/variables")
def update_variables(dataset_id: int, body: VariablesPatch, user: CurrentUser, db: DB):
    """Edits metadata only. The stored data is never modified; missing codes are applied when data is read."""
    v = latest_version(get_dataset(db, user, dataset_id))
    by_name = {x.column_name: x for x in v.variables}
    reprofile = False
    for name, patch in body.variables.items():
        var = by_name.get(name)
        if var is None:
            raise AppError("unknown_variable", f"Variable '{name}' not found.", 404)
        data = patch.model_dump(exclude_unset=True)
        if "level" in data:
            if data["level"] not in LEVELS:
                raise AppError("invalid_level", f"'{data['level']}' is not a valid measurement level.", 422)
            var.measurement_level = data["level"]
            reprofile = True
        if "label" in data:
            var.variable_label = (data["label"] or "").strip() or None
        if "category_labels" in data:
            var.category_labels_json = {str(k): str(val) for k, val in (data["category_labels"] or {}).items() if str(val).strip()}
        if "category_order" in data:
            var.category_order_json = data["category_order"] or []
        if "missing_codes" in data:
            var.missing_value_codes_json = data["missing_codes"] or []
            reprofile = True
        if "unit" in data:
            var.unit = data["unit"] or None
        if "scale_group" in data:
            var.scale_group = data["scale_group"] or None
    db.flush()
    if reprofile:
        db.refresh(v)
        refresh_profile(db, v)
    audit(db, user, "dataset.variables_update", v.dataset.project_id, dataset=dataset_id,
          variables=list(body.variables))
    db.commit()
    db.refresh(v)
    prof = {c["name"]: c for c in (v.profile_json or {}).get("variables", [])}
    return {"items": [variable_out(x, prof) for x in v.variables], "version_id": v.id}


@router.get("/datasets/{dataset_id}/download")
def download(dataset_id: int, user: CurrentUser, db: DB, format: str = Query("csv", pattern="^(csv|parquet|original|xlsx)$"),
             version_id: int | None = None):
    ds = get_dataset(db, user, dataset_id)
    v = _version(db, user, dataset_id, version_id)
    base = safe_filename(ds.name, "dataset")
    if format == "parquet":
        return FileResponse(resolve(v.data_path), filename=f"{base}.parquet", media_type="application/octet-stream")
    if format == "xlsx":
        return _xlsx_response(ds, v)
    if format == "original":
        if not ds.storage_path:
            raise AppError("not_available", "This dataset has no original upload (it was derived).", 404)
        return FileResponse(resolve(ds.storage_path), filename=ds.source_filename or base)

    def stream():
        for batch in pl.scan_parquet(resolve(v.data_path)).collect_batches(chunk_size=50000):
            buf = batch.write_csv(include_header=stream.first)
            stream.first = False
            yield buf.encode("utf-8")

    stream.first = True
    return StreamingResponse(stream(), media_type="text/csv",
                             headers={"Content-Disposition": f'attachment; filename="{base}.csv"'})


# --------------------------------------------------------------------------- transformations


class ScaleScoreIn(BaseModel):
    items: list[str] = Field(min_length=2)
    name: str = Field(min_length=1, max_length=120)
    method: str = Field("mean", pattern="^(mean|sum)$")
    reverse: list[str] = Field(default_factory=list)
    scale_min: float | None = None
    scale_max: float | None = None
    min_valid: int | None = Field(None, ge=1)


@router.post("/datasets/{dataset_id}/transforms/scale-score", status_code=201)
def scale_score(dataset_id: int, body: ScaleScoreIn, user: CurrentUser, db: DB):
    """Adds a computed scale score as a new dataset version; earlier versions stay unchanged."""
    ds = get_dataset(db, user, dataset_id)
    v = latest_version(ds)
    names = {x.column_name for x in v.variables}
    missing = [i for i in body.items if i not in names]
    if missing:
        raise AppError("unknown_variable", f"Unknown items: {', '.join(missing)}.", 422)
    new_name = body.name.strip()
    if new_name in names:
        raise AppError("duplicate_column_name", f"A variable named '{new_name}' already exists.", 409)
    df = store.load_frame(v, body.items)
    for c in body.items:
        df[c] = df[c].astype(float)
    for c in body.reverse:
        lo = body.scale_min if body.scale_min is not None else df[c].min()
        hi = body.scale_max if body.scale_max is not None else df[c].max()
        df[c] = lo + hi - df[c]
    valid = df[body.items].notna().sum(axis=1)
    need = body.min_valid or len(body.items)
    score = df[body.items].mean(axis=1) if body.method == "mean" else df[body.items].sum(axis=1, min_count=1)
    score[valid < need] = None
    full = pl.scan_parquet(resolve(v.data_path)).collect()
    full = full.with_columns(pl.Series(new_name, score.to_numpy(), dtype=pl.Float64))
    rel = new_relative_path("datasets", ".parquet")
    full.write_parquet(resolve(rel), compression="zstd")
    note = (f"Added '{new_name}' = {body.method} of {', '.join(body.items)}"
            + (f" (reversed: {', '.join(body.reverse)})" if body.reverse else "")
            + (f"; requires {need} answered items" if need else ""))
    nv = add_version(db, ds, rel, note=note, carry_from=v)
    nv_var = next(x for x in nv.variables if x.column_name == new_name)
    nv_var.measurement_level = "numeric"
    nv_var.variable_label = new_name
    nv_var.inference_note = note
    ds.lineage_json = {**(ds.lineage_json or {}), "transforms": [*(ds.lineage_json or {}).get("transforms", []),
                                                                 {"version": nv.version_number, "note": note}]}
    audit(db, user, "dataset.transform", ds.project_id, dataset=ds.id, note=note)
    db.commit()
    return dataset_out(ds, full=True)


# --------------------------------------------------------------------------- data cleaning


class CleanIn(BaseModel):
    steps: list[dict[str, Any]] = Field(default_factory=list, max_length=100)


def _raw_frame(v) -> pl.DataFrame:
    return pl.read_parquet(resolve(v.data_path))


@router.get("/datasets/{dataset_id}/cleaning/suggestions")
def cleaning_suggestions(dataset_id: int, user: CurrentUser, db: DB):
    """Problems found in the latest version, each with the step that fixes it."""
    v = latest_version(get_dataset(db, user, dataset_id))
    df = _raw_frame(v)
    levels = {x.column_name: x.measurement_level for x in v.variables}
    return {"version_id": v.id, "version_number": v.version_number, "summary": cleaning.summary(df),
            "suggestions": cleaning.suggest(df, levels)}


@router.post("/datasets/{dataset_id}/cleaning/preview")
def cleaning_preview(dataset_id: int, body: CleanIn, user: CurrentUser, db: DB):
    v = latest_version(get_dataset(db, user, dataset_id))
    try:
        return cleaning.preview(_raw_frame(v), body.steps)
    except cleaning.CleaningError as exc:
        raise AppError("invalid_cleaning_step", str(exc), 422) from exc


@router.post("/datasets/{dataset_id}/cleaning/apply", status_code=201)
def cleaning_apply(dataset_id: int, body: CleanIn, user: CurrentUser, db: DB):
    """Saves the cleaned data as a new version. Earlier versions (including the raw import) stay unchanged."""
    ds = get_dataset(db, user, dataset_id)
    v = latest_version(ds)
    if not body.steps:
        raise AppError("no_cleaning_steps", "Choose at least one cleaning step.", 422)
    try:
        df, log = cleaning.apply(_raw_frame(v), body.steps)
    except cleaning.CleaningError as exc:
        raise AppError("invalid_cleaning_step", str(exc), 422) from exc
    df = df.drop(cleaning.ROW)
    if df.height == 0:
        raise AppError("empty_result", "These steps would remove every row. Check the steps that remove rows.", 422)
    rel = new_relative_path("datasets", ".parquet")
    df.write_parquet(resolve(rel), compression="zstd")
    n = len(body.steps)
    note = f"Cleaned ({n} {'step' if n == 1 else 'steps'}) from version {v.version_number}"
    nv = add_version(db, ds, rel, note=note, carry_from=v)
    # Renamed columns keep their labels and levels; columns whose type changed are re-inferred.
    renamed: dict[str, str] = {}
    for s in body.steps:
        if s.get("op") == "rename" and s.get("column") and s.get("to"):
            src = next((k for k, cur in renamed.items() if cur == s["column"]), s["column"])
            renamed[src] = str(s["to"]).strip()
    old = {x.column_name: x for x in v.variables}
    source_of = {new: src for src, new in renamed.items()}
    for var in nv.variables:
        src = old.get(source_of.get(var.column_name, var.column_name))
        if src is None:
            continue
        if var.column_name not in old:
            for f in VAR_FIELDS:
                setattr(var, f, getattr(src, f))
        if src.data_type != var.data_type:
            var.measurement_level = var.inferred_level
    db.flush()
    db.refresh(nv)
    refresh_profile(db, nv)
    entry = {"version": nv.version_number, "from_version": v.version_number, "steps": body.steps,
             "log": [x["message"] for x in log], "rows_before": v.row_count, "rows_after": nv.row_count}
    lineage = ds.lineage_json or {}
    ds.lineage_json = {**lineage, "cleaning": [*lineage.get("cleaning", []), entry],
                       "transforms": [*lineage.get("transforms", []), {"version": nv.version_number, "note": note}]}
    audit(db, user, "dataset.clean", ds.project_id, dataset=ds.id, steps=n)
    db.commit()
    return {"dataset": dataset_out(ds, full=True), "log": log}


def _xlsx_response(ds, v) -> StreamingResponse:
    """The data, plus a 'Cleaning log' sheet for cleaned versions that can be cited in the methodology."""
    import io

    from openpyxl import Workbook

    frame = pl.read_parquet(resolve(v.data_path))
    if frame.height > 1_048_000:
        raise AppError("too_large_for_excel", "This dataset has more rows than Excel allows. Download CSV instead.", 422)
    wb = Workbook(write_only=True)
    data = wb.create_sheet("Data")
    data.append(frame.columns)
    for row in frame.iter_rows():
        data.append([x.isoformat() if hasattr(x, "isoformat") else x for x in row])
    entries = [e for e in (ds.lineage_json or {}).get("cleaning", []) if e["version"] <= v.version_number]
    if entries:
        logs = wb.create_sheet("Cleaning log")
        first = min(ds.versions, key=lambda x: x.version_number)
        logs.append([f"Data cleaning log — {ds.name}"])
        logs.append([f"Raw data: version {first.version_number}"
                     + (f" ({ds.source_filename})" if ds.source_filename else "") + f", {first.row_count:,} rows"])
        logs.append([f"This file: version {v.version_number}, {v.row_count:,} rows"])
        logs.append([])
        logs.append(["Version", "Step", "What was done"])
        for e in entries:
            for i, msg in enumerate(e["log"], start=1):
                logs.append([e["version"], i, msg])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = safe_filename(f"{ds.name}_v{v.version_number}", "dataset")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})
