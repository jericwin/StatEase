"""Data cleaning: suggested fixes, preview, saving a cleaned version, and exporting it with the cleaning log."""

from __future__ import annotations

import io

import polars as pl
import pytest
from openpyxl import load_workbook

from app.data import cleaning

API = "/api/v1"

MESSY = """ID,Gender,Age,Q1,Q2,Notes
1,Male,21,4,3,
2,male ,22,5,4,
3,Female,"1,000",3,4,
4,F,n/a,99,5,
5,M,24,5,2,
5,M,24,5,2,
6,N/A,abc,4,3,
,,,,,
7,female,19,55,4,
8,Female,20,2,5,
"""


def test_engine_fixes_typical_survey_problems():
    raw = pl.read_csv(io.StringIO(MESSY), infer_schema_length=0).with_columns(pl.col("Q1", "Q2").cast(pl.Int64))
    sug = cleaning.suggest(raw, {"ID": "identifier"})
    kinds = {s["kind"] for s in sug}
    assert {"empty_rows", "empty_columns", "whitespace", "missing_tokens", "categories", "numbers_as_text",
            "likert_range", "duplicates"} <= kinds, kinds
    df, log = cleaning.apply(raw, [s["step"] for s in sug if s["default"]])
    df = df.drop(cleaning.ROW)
    assert "Notes" not in df.columns
    assert df.height == 8  # empty row and one duplicate removed
    assert sorted(df["Gender"].drop_nulls().unique().to_list()) == ["Female", "Male"]
    assert df["Age"].dtype == pl.Int64 and df["Age"].null_count() == 2  # "n/a" and "abc"
    assert df["Q1"].drop_nulls().max() == 5  # 55 and 99 are not on a 1–5 scale
    assert all(entry["message"] for entry in log)


def test_outliers_are_never_suggested_and_steps_validate():
    raw = pl.DataFrame({"Score": [10, 11, 12, 11, 10, 12, 11, 95]})
    assert not [s for s in cleaning.suggest(raw) if "outlier" in s["title"].lower()]
    df, log = cleaning.apply(raw, [{"op": "outliers", "column": "Score", "method": "iqr", "action": "missing"}])
    assert df["Score"].null_count() == 1 and "IQR" in log[0]["message"]
    with pytest.raises(cleaning.CleaningError, match="Step 1: Column not found"):
        cleaning.apply(raw, [{"op": "rename", "column": "Nope", "to": "X"}])


def test_cleaning_api_saves_new_version_and_exports_log(client):
    p = client.post(f"{API}/projects", json={"name": "Cleaning study"}).json()
    up = client.post(f"{API}/datasets/uploads", files={"file": ("messy.csv", MESSY.encode(), "text/csv")}).json()
    ds = client.post(f"{API}/datasets", json={"project_id": p["id"], "upload_id": up["upload"]["upload_id"]}).json()
    did = ds["id"]

    sug = client.get(f"{API}/datasets/{did}/cleaning/suggestions").json()
    assert sug["summary"]["empty_rows"] >= 1 and sug["suggestions"]
    steps = [s["step"] for s in sug["suggestions"] if s["default"]]
    steps.append({"op": "rename", "column": "Q1", "to": "Item 1"})

    pv = client.post(f"{API}/datasets/{did}/cleaning/preview", json={"steps": steps}).json()
    assert pv["after"]["rows"] < pv["before"]["rows"] and pv["changed"] and "Item 1" in pv["columns"]
    bad = client.post(f"{API}/datasets/{did}/cleaning/preview", json={"steps": [{"op": "fill_missing", "column": "Gone"}]})
    assert bad.status_code == 422 and "Gone" in bad.json()["error"]["message"]

    r = client.post(f"{API}/datasets/{did}/cleaning/apply", json={"steps": steps})
    assert r.status_code == 201, r.text
    out = r.json()["dataset"]
    assert out["version"]["number"] == 2 and [v["number"] for v in out["versions"]] == [1, 2]
    assert out["lineage"]["cleaning"][0]["log"]
    names = {v["name"] for v in out["variables"]}
    assert "Item 1" in names and "Q1" not in names

    # Raw data is untouched as version 1.
    v1 = out["versions"][0]["id"]
    raw_csv = client.get(f"{API}/datasets/{did}/download", params={"format": "csv", "version_id": v1}).text
    assert "male " in raw_csv
    # Cleaned export: data sheet plus the cleaning log.
    x = client.get(f"{API}/datasets/{did}/download", params={"format": "xlsx"})
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Data", "Cleaning log"]
    log_text = " ".join(str(c.value) for row in wb["Cleaning log"].iter_rows() for c in row if c.value)
    assert "Raw data: version 1" in log_text and "Renamed Q1 to Item 1" in log_text
    # Removing everything is refused.
    nuke = client.post(f"{API}/datasets/{did}/cleaning/apply", json={"steps": [{"op": "drop_missing", "columns": ["Notes"]}]})
    assert nuke.status_code in (422,)
