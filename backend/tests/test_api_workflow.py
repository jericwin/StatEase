"""End-to-end API workflow: import -> recommendations -> analyses -> results -> DOCX/XLSX export."""

from __future__ import annotations

import io

import numpy as np
import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook

API = "/api/v1"


def survey_csv(n: int = 120) -> bytes:
    rng = np.random.default_rng(5)
    df = pd.DataFrame({
        "Respondent ID": np.arange(1, n + 1),
        "Gender": rng.choice(["Male", "Female"], n),
        "Education": rng.choice(["High school", "College", "Graduate"], n),
        "Age": rng.integers(18, 60, n),
        "Service Quality": np.round(rng.normal(3.6, 0.6, n), 2),
        "Customer Satisfaction": np.round(rng.normal(3.8, 0.7, n), 2),
        "Purchase Intention": np.round(rng.normal(3.5, 0.8, n), 2),
        "SQ1": rng.integers(1, 6, n), "SQ2": rng.integers(1, 6, n), "SQ3": rng.integers(1, 6, n),
    })
    df.loc[3, "Customer Satisfaction"] = np.nan
    df.loc[7, "Age"] = np.nan
    return df.to_csv(index=False).encode()


@pytest.fixture(scope="module")
def project(client):
    r = client.post(f"{API}/projects", json={"name": "Workflow test"})
    assert r.status_code == 201, r.text
    return r.json()


@pytest.fixture(scope="module")
def dataset(client, project):
    up = client.post(f"{API}/datasets/uploads", files={"file": ("survey.csv", survey_csv(), "text/csv")})
    assert up.status_code == 201, up.text
    body = up.json()
    levels = {v["column_name"]: v["measurement_level"] for v in body["variables"]}
    assert levels["Respondent ID"] == "identifier"
    assert levels["Gender"] == "binary"
    assert levels["Education"] == "nominal"
    assert levels["Customer Satisfaction"] == "numeric"
    assert levels["SQ1"] == "ordinal"
    assert body["detected"]["delimiter"] == ","
    r = client.post(f"{API}/datasets", json={"project_id": project["id"], "upload_id": body["upload"]["upload_id"],
                                             "name": "Survey",
                                             "variables": {"Education": {"measurement_level": "ordinal",
                                                                         "category_order": ["High school", "College",
                                                                                            "Graduate"]}}})
    assert r.status_code == 201, r.text
    ds = r.json()
    assert ds["row_count"] == 120
    assert ds["profile"]["total_missing"] == 2
    edu = next(v for v in ds["variables"] if v["name"] == "Education")
    assert edu["level"] == "ordinal" and edu["category_order"][0] == "High school"
    return ds


def test_preview_paging_and_filter(client, dataset):
    r = client.get(f"{API}/datasets/{dataset['id']}/preview", params={"limit": 10, "sort": "Age", "desc": True})
    assert r.status_code == 200
    page = r.json()
    assert page["total"] == 120 and len(page["rows"]) == 10
    ages = [row[page["columns"].index("Age")] for row in page["rows"]]
    assert ages == sorted(ages, reverse=True)
    r = client.get(f"{API}/datasets/{dataset['id']}/preview",
                   params={"filters": '[{"column": "Gender", "op": "eq", "value": "Male"}]'})
    assert all(row[page["columns"].index("Gender")] == "Male" for row in r.json()["rows"])


def test_recommendations(client, dataset):
    r = client.get(f"{API}/analyses/recommendations", params={"dataset_id": dataset["id"]})
    assert r.status_code == 200
    recs = r.json()["recommendations"]
    types = {x["analysis_type"] for x in recs}
    assert "descriptives" in types and "correlation" in types and "reliability" in types
    titles = " ".join(x["title"] for x in recs)
    assert "Gender" in titles
    gender_cmp = [x for x in recs if x["config"]["variables"].get("group") == "Gender"]
    assert gender_cmp and gender_cmp[0]["analysis_type"] == "compare_two_groups"
    assert gender_cmp[0]["title"] in ("Compare Female and Male", "Compare Male and Female")
    assert not any(x["config"]["variables"].get("dependent") == "Education" for x in recs)
    assert not any("Respondent ID" in str(x["config"]) for x in recs)


def test_full_analysis_and_export(client, project, dataset):
    vid = dataset["version"]["id"]
    configs = []
    for atype, variables in [
        ("descriptives", {"variables": ["Age", "Customer Satisfaction"]}),
        ("independent_t", {"dependent": "Customer Satisfaction", "group": "Gender"}),
        ("one_way_anova", {"dependent": "Purchase Intention", "factor": "Education"}),
        ("correlation", {"variables": ["Service Quality", "Customer Satisfaction", "Purchase Intention"]}),
    ]:
        r = client.post(f"{API}/analyses/configurations", json={"project_id": project["id"], "dataset_version_id": vid,
                                                                "analysis_type": atype, "variables": variables})
        assert r.status_code == 201, r.text
        assert r.json()["problems"] == []
        configs.append(r.json()["id"])

    bad = client.post(f"{API}/analyses/configurations", json={
        "project_id": project["id"], "dataset_version_id": vid, "analysis_type": "independent_t",
        "variables": {"dependent": "Gender", "group": "Age"}}).json()
    assert any("needs a" in p for p in bad["problems"])

    r = client.post(f"{API}/analyses/runs", json={"configuration_ids": configs, "wait": True})
    assert r.status_code == 202
    runs = r.json()["items"]
    assert all(x["status"] == "completed" for x in runs), runs

    res = client.get(f"{API}/analyses/runs/{runs[1]['id']}/results").json()["result"]
    main = next(t for t in res["tables"] if t["id"] == "main")
    assert main["rows"][0]["df"] > 0 and 0 <= main["rows"][0]["p"] <= 1
    assert res["missing"]["excluded"] == 1
    assert "stats.ttest_ind" in res["code"]

    r = client.put(f"{API}/analyses/runs/{runs[1]['id']}/interpretation", json={"text": "Edited interpretation."})
    assert r.status_code == 200

    opts = {"title": "Survey Report", "run_ids": [x["id"] for x in runs]}
    r = client.post(f"{API}/exports", json={"project_id": project["id"], "format": "docx", "options": opts})
    assert r.status_code == 201, r.text
    exp = r.json()
    assert exp["status"] == "completed" and exp["size_bytes"] > 0
    blob = client.get(exp["download_url"]).content
    doc = Document(io.BytesIO(blob))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Survey Report" in text and "Edited interpretation." in text
    assert len(doc.tables) >= 5
    assert len(doc.inline_shapes) >= 1
    cells = {c.text for t in doc.tables for row in t.rows for c in row.cells}
    assert any(c.startswith(".") or c == "< .001" for c in cells)  # APA p-values

    # Excluding sections must remove them from the document.
    opts2 = {**opts, "sections": {"title": True, "main": True}, "inferential": {"p": False}}
    exp2 = client.post(f"{API}/exports", json={"project_id": project["id"], "format": "docx", "options": opts2}).json()
    doc2 = Document(io.BytesIO(client.get(exp2["download_url"]).content))
    headers = {c.text for t in doc2.tables for c in t.rows[0].cells}
    assert "p" not in headers and "Interpretation" not in "\n".join(p.text for p in doc2.paragraphs)
    assert len(doc2.inline_shapes) == 0

    r = client.post(f"{API}/exports", json={"project_id": project["id"], "format": "xlsx",
                                            "options": {**opts, "include_data": True}})
    assert r.status_code == 201, r.text
    wb = load_workbook(io.BytesIO(client.get(r.json()["download_url"]).content))
    assert {"Overview", "Descriptive Statistics", "T-Test Results", "ANOVA Results", "Correlations",
            "Variable Dictionary", "Data"} <= set(wb.sheetnames)
    ws = wb["T-Test Results"]
    numeric = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, float)]
    assert numeric, "statistics must be stored as numbers"
    assert wb["Data"].max_row == 121


def test_persistence_after_restart(client, project):
    """A new app instance (fresh engine/session) sees the saved configurations and results."""
    from fastapi.testclient import TestClient

    from app.main import create_app

    with TestClient(create_app()) as c2:
        items = c2.get(f"{API}/analyses/results", params={"project_id": project["id"]}).json()["items"]
        assert len(items) >= 4
        assert all(i["result"] for i in items if i["run"]["status"] == "completed")


def test_scale_score_creates_new_version(client, dataset):
    r = client.post(f"{API}/datasets/{dataset['id']}/transforms/scale-score",
                    json={"items": ["SQ1", "SQ2", "SQ3"], "name": "SQ_mean"})
    assert r.status_code == 201, r.text
    ds = r.json()
    assert ds["version"]["number"] == 2
    assert any(v["name"] == "SQ_mean" and v["level"] == "numeric" for v in ds["variables"])
    edu = next(v for v in ds["variables"] if v["name"] == "Education")
    assert edu["level"] == "ordinal"  # metadata carried to the new version


def test_variable_edit_and_missing_codes(client, project):
    csv = b"score,group\n1,a\n2,a\n99,b\n4,b\n5,b\n"
    up = client.post(f"{API}/datasets/uploads", files={"file": ("m.csv", csv, "text/csv")}).json()
    ds = client.post(f"{API}/datasets", json={"project_id": project["id"], "upload_id": up["upload"]["upload_id"]}).json()
    r = client.patch(f"{API}/datasets/{ds['id']}/variables",
                     json={"variables": {"score": {"missing_codes": [99], "level": "numeric", "label": "Test score"}}})
    assert r.status_code == 200
    score = next(v for v in r.json()["items"] if v["name"] == "score")
    assert score["profile"]["missing"] == 1 and score["label"] == "Test score"


def test_malformed_and_unsupported_files(client):
    r = client.post(f"{API}/datasets/uploads", files={"file": ("x.pdf", b"%PDF-1", "application/pdf")})
    assert r.status_code == 415 and r.json()["error"]["code"] == "unsupported_format"
    r = client.post(f"{API}/datasets/uploads", files={"file": ("e.csv", b"", "text/csv")})
    assert r.status_code == 400 and r.json()["error"]["code"] == "empty_file"
    ragged = b"a,b,c\n1,2,3\n4,5\n6,7,8,9\n" + b"1,2,3\n" * 5
    r = client.post(f"{API}/datasets/uploads", files={"file": ("r.csv", ragged, "text/csv")})
    assert r.status_code == 201
    assert any(i["code"] == "inconsistent_rows" for i in r.json()["issues"])
    dup = b"x,x,,y\n1,2,3,4\n5,6,7,8\n"
    r = client.post(f"{API}/datasets/uploads", files={"file": ("d.csv", dup, "text/csv")}).json()
    assert r["columns"] == ["x", "x_2", "column_3", "y"]
    codes = {i["code"] for i in r["issues"]}
    assert {"duplicate_column_name", "empty_column_name"} <= codes


def test_excel_import_with_sheets(client, project):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Notes"
    wb.active.append(["just notes"])
    ws = wb.create_sheet("Data")
    ws.append(["name", "score"])
    for i in range(10):
        ws.append([f"p{i}", i * 1.5])
    buf = io.BytesIO()
    wb.save(buf)
    up = client.post(f"{API}/datasets/uploads", files={"file": ("book.xlsx", buf.getvalue(), "application/octet-stream")})
    assert up.status_code == 201, up.text
    assert up.json()["detected"]["sheets"] == ["Notes", "Data"]
    pv = client.post(f"{API}/datasets/uploads/{up.json()['upload']['upload_id']}/preview",
                     json={"options": {"sheet": "Data"}}).json()
    assert pv["columns"] == ["name", "score"] and len(pv["rows"]) == 10


def test_sql_workbench_read_only_and_save(client, project, dataset):
    r = client.get(f"{API}/sql/schema", params={"project_id": project["id"]})
    assert r.status_code == 200
    table = next(t["name"] for t in r.json()["tables"] if t["dataset_id"] == dataset["id"])
    q = f'SELECT "Gender", COUNT(*) AS n, AVG("Age") AS mean_age FROM {table} GROUP BY 1 ORDER BY 1'
    r = client.post(f"{API}/sql/execute", json={"sql": q, "project_id": project["id"]})
    assert r.status_code == 200, r.text
    assert r.json()["row_count"] == 2
    for bad in (f"DROP VIEW {table}", f"DELETE FROM {table}", "COPY (SELECT 1) TO 'x.csv'",
                "SELECT * FROM read_csv('C:/Windows/win.ini')"):
        r = client.post(f"{API}/sql/execute", json={"sql": bad, "project_id": project["id"]})
        assert r.status_code in (403, 422), (bad, r.text)
    r = client.post(f"{API}/sql/save-dataset", json={"sql": f'SELECT * FROM {table} WHERE "Age" >= 30',
                                                     "project_id": project["id"], "name": "Age 30+"})
    assert r.status_code == 201, r.text
    derived = r.json()
    current = client.get(f"{API}/datasets/{dataset['id']}").json()["version"]["id"]
    assert derived["lineage"]["sql"] and derived["parent_version_id"] == current
    r = client.post(f"{API}/sql/export", json={"sql": f"SELECT * FROM {table}", "project_id": project["id"],
                                               "format": "xlsx"})
    assert r.status_code == 200 and len(r.content) > 0


def test_sql_sample_database(client):
    r = client.get(f"{API}/sql/samples")
    assert r.status_code == 200
    cid = r.json()["connection"]["id"]
    ex = next(e for e in r.json()["examples"] if e["title"].startswith("Window"))
    r = client.post(f"{API}/sql/execute", json={"sql": ex["sql"], "connection_id": cid})
    assert r.status_code == 200, r.text
    assert "price_rank" in [c["name"] for c in r.json()["columns"]]
    r = client.post(f"{API}/sql/execute", json={"sql": "DELETE FROM customers", "connection_id": cid})
    assert r.status_code == 403
    schema = client.get(f"{API}/sql/schema", params={"connection_id": cid}).json()
    items = next(t for t in schema["tables"] if t["name"] == "order_items")
    assert items["foreign_keys"]


def test_python_console_sandbox(client, dataset):
    ok = client.post(f"{API}/python/run", json={"code": "show(df.describe())\nprint(len(df))",
                                                  "dataset_id": dataset["id"]}).json()
    assert ok["status"] == "completed", ok
    assert "120" in ok["stdout"] and ok["tables"]
    blocked = client.post(f"{API}/python/run", json={"code": "import socket"}).json()
    assert blocked["status"] == "error" and "not allowed" in blocked["stderr"]
    net = "import urllib.request\nurllib.request.urlopen('http://example.com', timeout=3)"
    blocked = client.post(f"{API}/python/run", json={"code": net}).json()
    assert blocked["status"] == "error" and "not allowed" in blocked["stderr"]
    slow = client.post(f"{API}/python/run", json={"code": "while True: pass"}).json()
    assert slow["status"] == "killed" and "time limit" in slow["stderr"]
    blocked = client.post(f"{API}/python/run", json={"code": "open('C:/Windows/win.ini').read()"}).json()
    assert blocked["status"] == "error" and "not allowed" in blocked["stderr"]


def test_demo_project(client):
    r = client.post(f"{API}/demo")
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    assert r.json()["report_id"]
    items = client.get(f"{API}/analyses/results", params={"project_id": pid}).json()["items"]
    # Fertilizer trial: growth by treatment, difference among the three treatments, and the best treatment.
    assert sorted(i["result"]["analysis_type"] for i in items) == ["compare_groups", "descriptives", "rank_groups"], items
    assert all(i["run"]["status"] == "completed" for i in items), [i["run"] for i in items]
    # Sub-items "1.1 Plant height" ... match columns that carry units, e.g. "Plant Height (cm)".
    desc = next(i["result"] for i in items if i["result"]["analysis_type"] == "descriptives")
    measures = {r.get("variable") or r.get("measure") for t in desc["tables"] for r in t["rows"]}
    assert {"Plant Height (cm)", "Number of Leaves", "Leaf Width (cm)", "Fresh Weight (g)"} <= measures, measures


def test_project_lifecycle(client):
    p = client.post(f"{API}/projects", json={"name": "Temp"}).json()
    assert client.patch(f"{API}/projects/{p['id']}", json={"name": "Renamed", "status": "archived"}).json()["status"] == "archived"
    assert client.post(f"{API}/projects/{p['id']}/duplicate").status_code == 201
    assert client.delete(f"{API}/projects/{p['id']}").status_code == 204
    assert client.get(f"{API}/projects/{p['id']}").status_code == 404
