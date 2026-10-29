"""Word report generation with python-docx: editable tables, numbered tables and figures, captions."""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import charts
from .options import ReportOptions, chosen_charts, filter_table, fmt_value, interpretation_text


def _border(cell, **edges: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, spec in edges.items():
        el = OxmlElement(f"w:{edge}")
        if spec == "none":
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), spec)
            el.set(qn("w:color"), "000000")
        borders.append(el)


def _page_number(paragraph) -> None:
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


class DocxReport:
    def __init__(self, opts: ReportOptions):
        self.opts = opts
        self.doc = Document()
        self.table_no = 0
        self.figure_no = 0
        apa = opts.style == "apa"
        self.font = "Times New Roman" if apa else "Calibri"
        normal = self.doc.styles["Normal"]
        normal.font.name = self.font
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), self.font)
        normal.font.size = Pt(12 if apa else 11)
        normal.paragraph_format.space_after = Pt(6)
        for lvl, size in ((1, 14), (2, 12), (3, 12)):
            st = self.doc.styles[f"Heading {lvl}"]
            st.font.name = self.font
            st.element.rPr.rFonts.set(qn("w:eastAsia"), self.font)
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.italic = lvl == 3 and apa
            st.font.color.rgb = RGBColor(0, 0, 0) if apa else RGBColor(0x1F, 0x3A, 0x5F)
        sec = self.doc.sections[0]
        sec.orientation = WD_ORIENT.PORTRAIT
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _page_number(footer)

    @property
    def width(self) -> float:
        return 6.5

    def para(self, text: str, *, italic=False, bold=False, size: float | None = None, align=None, style=None):
        p = self.doc.add_paragraph(style=style)
        run = p.add_run(text)
        run.italic, run.bold = italic, bold
        if size:
            run.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        return p

    def table(self, t: dict[str, Any]) -> None:
        cols = t["columns"]
        if not cols:
            return
        self.table_no += 1
        apa = self.opts.style == "apa"
        cap = self.doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(12)
        r = cap.add_run(f"Table {self.table_no}")
        r.bold = True
        if apa:
            cap2 = self.doc.add_paragraph()
            cap2.paragraph_format.keep_with_next = True
            cap2.add_run(t["title"]).italic = True
        else:
            cap.add_run(f". {t['title']}")
        tbl = self.doc.add_table(rows=1, cols=len(cols))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True
        size = Pt(10 if len(cols) <= 7 else 8.5 if len(cols) <= 10 else 7.5)
        hdr = tbl.rows[0].cells
        for i, c in enumerate(cols):
            hdr[i].text = ""
            run = hdr[i].paragraphs[0].add_run(c["label"])
            run.font.size = size
            run.bold = not apa
            hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if c["fmt"] == "text" else WD_ALIGN_PARAGRAPH.CENTER
            _border(hdr[i], top="8", bottom="6", left="none", right="none")
        rows = t["rows"]
        for ri, row in enumerate(rows):
            cells = tbl.add_row().cells
            for i, c in enumerate(cols):
                v = row.get(c["key"])
                val = fmt_value(v, "text" if isinstance(v, str) else c["fmt"], self.opts)
                if c["key"] in (row.get("_manual") or []) and val:
                    val += "†"
                if row.get(f"{c['key']}_sig"):
                    val += row[f"{c['key']}_sig"]
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                run = p.add_run(val)
                run.font.size = size
                if row.get("_style") == "total":
                    run.bold = True
                if row.get("_style") == "muted":
                    run.italic = True
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c["fmt"] == "text" else WD_ALIGN_PARAGRAPH.RIGHT
                p.paragraph_format.space_after = Pt(0)
                last = ri == len(rows) - 1
                _border(cells[i], top="none", bottom="8" if last else "none", left="none", right="none")
        for cell in hdr:
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        notes = [n for n in t.get("notes") or [] if n]
        if notes:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            r = p.add_run("Note. ")
            r.italic = True
            r.font.size = Pt(9.5)
            r2 = p.add_run(" ".join(notes))
            r2.font.size = Pt(9.5)

    def figure(self, spec: dict[str, Any]) -> bool:
        try:
            png = charts.render(spec)
        except Exception:  # noqa: BLE001 - never claim a figure that was not generated
            return False
        self.figure_no += 1
        apa = self.opts.style == "apa"
        cap = self.doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(12)
        cap.add_run(f"Figure {self.figure_no}").bold = True
        if apa:
            c2 = self.doc.add_paragraph()
            c2.paragraph_format.keep_with_next = True
            c2.add_run(spec["title"]).italic = True
        else:
            cap.add_run(f". {spec['title']}")
        self.doc.add_picture(io.BytesIO(png), width=Inches(5.8))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = " ".join(x for x in (spec.get("caption"), spec.get("note")) if x)
        if caption:
            p = self.doc.add_paragraph()
            r = p.add_run("Note. ")
            r.italic = True
            r.font.size = Pt(9.5)
            p.add_run(caption).font.size = Pt(9.5)
        return True


def build_docx(path: Path, opts: ReportOptions, info: dict[str, Any], items: list[dict[str, Any]]) -> dict[str, Any]:
    """items: [{run_id, result, edited_interpretation}] in report order. Returns counts of what was written."""
    rep = DocxReport(opts)
    doc = rep.doc
    if opts.on("title"):
        p = rep.para(opts.title, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_after = Pt(4)
        if opts.subtitle:
            rep.para(opts.subtitle, align=WD_ALIGN_PARAGRAPH.CENTER)
        meta = [x for x in (opts.author, info.get("project")) if x]
        if opts.on("date"):
            meta.append(datetime.now().strftime("%B %d, %Y"))
        if meta:
            rep.para(" · ".join(meta), align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    elif opts.on("date"):
        rep.para(f"Generated {datetime.now():%B %d, %Y}", size=10)

    if opts.on("dataset_description"):
        doc.add_heading("Data", level=1)
        lineage = info.get("lineage")
        rep.para(f"Dataset: {info['dataset']} ({info['rows']:,} rows × {info['columns']} columns)"
                 + (f", imported from {info['source']}" if info.get("source") else "")
                 + (f". Derived from another dataset by: {lineage}" if lineage else "") + ".")
    if opts.on("alpha"):
        rep.para(f"Statistical significance was evaluated at α = {info['alpha']:g}"
                 f"{' (two-sided unless stated otherwise)'}. p-values below .001 are reported as p < .001.")

    if opts.on("variables") and info.get("variables"):
        doc.add_heading("Research Variables", level=1)
        used = {v for it in items for vv in (it["result"].get("variables") or {}).values()
                for v in ([vv] if isinstance(vv, str) else vv or [])}
        rows = [{"name": v["name"], "label": v.get("label") or "", "level": v["level"]}
                for v in info["variables"] if v["name"] in used]
        if rows:
            rep.table({"title": "Variables Used in the Analyses", "kind": "main", "notes": [],
                       "columns": [{"key": "name", "label": "Variable", "fmt": "text"},
                                   {"key": "label", "label": "Label", "fmt": "text"},
                                   {"key": "level", "label": "Measurement level", "fmt": "text"}], "rows": rows})

    tables = figures = skipped_figures = 0
    for idx, it in enumerate(items):
        res = it["result"]
        head = res.get("title") or "Analysis"
        doc.add_heading(f"{head}: {res['heading']}" if res.get("heading") else head, level=1)
        if res.get("plain_title"):
            rep.para(res["plain_title"], italic=True, size=10)
        if opts.on("methods") and res.get("method", {}).get("description"):
            doc.add_heading("Method", level=3)
            m = res["method"]
            text = m["description"]
            if m.get("missing_data"):
                text += " Missing data: " + m["missing_data"]
            rep.para(text)
        for n in res.get("notes") or []:
            if opts.on("methods") or "excluded" in n:
                rep.para(n, size=10)
        for t in res.get("tables") or []:
            ft = filter_table(t, opts)
            if ft and ft["rows"]:
                rep.table(ft)
                tables += 1
        if opts.on("assumptions") and res.get("assumptions"):
            doc.add_heading("Assumption Checks", level=3)
            for a in res["assumptions"]:
                p = doc.add_paragraph(style="List Bullet")
                status = {"ok": "Met", "warning": "Caution", "violated": "Not met", "info": "Note"}.get(a["status"], "")
                p.add_run(f"{a['name']} — {status}. ").bold = True
                p.add_run(a["message"])
        if res.get("warnings") and (opts.on("assumptions") or opts.on("interpretation")):
            for w in res["warnings"]:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run("Warning: ").bold = True
                p.add_run(w)
        for spec in chosen_charts(res, it["run_id"], opts):
            if rep.figure(spec):
                figures += 1
            else:
                skipped_figures += 1
        if opts.on("interpretation"):
            text = interpretation_text(res, it["run_id"], opts, it.get("edited_interpretation"))
            if text:
                doc.add_heading("Interpretation", level=3)
                for block in text.split("\n\n"):
                    rep.para(block)
        if opts.on("references") and res.get("method", {}).get("references"):
            doc.add_heading("References", level=3)
            for ref in res["method"]["references"]:
                rep.para(ref, size=10)
        if opts.on("code") and res.get("code"):
            doc.add_heading("Reproducible Code", level=3)
            p = doc.add_paragraph()
            r = p.add_run(res["code"])
            r.font.name = "Consolas"
            r.font.size = Pt(8)
        if idx < len(items) - 1 and (tables > 3 or figures):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.core_properties.title = opts.title
    doc.core_properties.author = opts.author or "StatEase Analytics"
    doc.save(path)
    return {"tables": tables, "figures": figures, "figures_failed": skipped_figures}


def verify_docx(path: Path) -> dict[str, Any]:
    d = Document(str(path))
    return {"paragraphs": len(d.paragraphs), "tables": len(d.tables), "images": len(d.inline_shapes)}


def build_report_docx(path: Path, report: dict[str, Any], info: dict[str, Any]) -> dict[str, Any]:
    """Word document from an edited report (resolved blocks, in the user's order)."""
    o = report["options"]
    opts = ReportOptions(title=report["title"], style=o.get("style", "apa"), decimals=int(o.get("decimals", 2)),
                         author=o.get("author"), subtitle=o.get("subtitle"))
    rep = DocxReport(opts)
    doc = rep.doc
    rep.para(report["title"], bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    if opts.subtitle:
        rep.para(opts.subtitle, align=WD_ALIGN_PARAGRAPH.CENTER)
    meta = [x for x in (opts.author, info.get("project"), datetime.now().strftime("%B %d, %Y")) if x]
    rep.para(" \u00b7 ".join(meta), align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    if info.get("synthetic"):
        rep.para("Note: this report uses fictional sample data for demonstration.", italic=True, size=10,
                 align=WD_ALIGN_PARAGRAPH.CENTER)
    objectives = [x for x in report.get("objectives") or [] if x.get("text")]
    if objectives:
        doc.add_heading("Objectives of the Study", level=1)
        for i, ob in enumerate(objectives, start=1):
            doc.add_paragraph(f"{i}. {ob['text']}")
    tables = figures = failed = 0
    methods_done: set[int] = set()
    blocks = [b for b in report["blocks"] if b.get("include") and not b.get("missing")]
    for idx, b in enumerate(blocks):
        t = b["type"]
        if t == "heading":
            doc.add_heading(b.get("resolved_text") or b.get("text") or "", level=1)
        elif t == "table":
            rep.table(b["resolved"])
            tables += 1
        elif t == "chart":
            if rep.figure(dict(b["resolved"])):
                figures += 1
            else:
                failed += 1
        elif t == "text":
            text = (b.get("resolved_text") or "").strip()
            if text:
                if b.get("title") and b.get("title") != "Interpretation":
                    doc.add_heading(b["title"], level=3)
                for part in text.split("\n\n"):
                    rep.para(part)
        run_id = (b.get("source") or {}).get("run_id")
        nxt = blocks[idx + 1] if idx + 1 < len(blocks) else None
        section_ends = nxt is None or nxt["type"] == "heading"
        if o.get("include_methods", True) and section_ends:
            for prev in blocks[: idx + 1][::-1]:
                rid = (prev.get("source") or {}).get("run_id")
                if prev["type"] == "heading":
                    break
                if rid and rid not in methods_done:
                    a = prev.get("analysis") or {}
                    m = a.get("method") or {}
                    if m.get("description"):
                        p = doc.add_paragraph()
                        r = p.add_run("Method. ")
                        r.italic = True
                        r.font.size = Pt(9.5)
                        p.add_run(f"{a.get('title')}: {m['description']}").font.size = Pt(9.5)
                    methods_done.add(rid)
        del run_id
    doc.core_properties.title = report["title"]
    doc.core_properties.author = opts.author or "StatEase Analytics"
    doc.save(path)
    return {"tables": tables, "figures": figures, "figures_failed": failed}
